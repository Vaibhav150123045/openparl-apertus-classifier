"""
Create train/val/test datasets for parliamentary affair topic classification.

Merges:
  - che_affairs_keywords.csv  → affair metadata + Curia Vista topic labels
  - che_affair_texts.ndjson.gz → affair texts (title, submitted text, reasoning, answer) in DE/FR/IT

Outputs:
  - Core dataset of 2,000 entries saved as a single multi-sheet Excel file (.xlsx)
    with tabs: 'Core_Dataset', 'Train', 'Validation', 'Test', 'Summary_Stats'
    Saved at: data/core_dataset_2000.xlsx (and data/splits/core_dataset_2000.xlsx)
  - Split CSV files in data/splits/:
    train.csv (1,400), val.csv (300), test.csv (300), combined_2000.csv (2,000)

Target variables: topic_ids, topics_de, topics_fr, topics_it (multi-label, pipe-separated)
"""

import os
import sys
import re
import argparse
from pathlib import Path
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Paths ────────────────────────────────────────────────────────────────────
PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
OUTPUT_DIR = DATA_DIR / "splits"
OUTPUT_DIR.mkdir(exist_ok=True)

KEYWORDS_CSV = DATA_DIR / "che_affairs_keywords.csv"
TEXTS_NDJSON = DATA_DIR / "che_affair_texts.ndjson.gz"
EXCEL_OUTPUT_PATH = DATA_DIR / "core_dataset_2000.xlsx"
EXCEL_SPLITS_PATH = OUTPUT_DIR / "core_dataset_2000.xlsx"

DEFAULT_SEED = 42
DEFAULT_SAMPLES = 2000


def sanitize_text_for_excel(val):
    """Clean control characters and truncate excessively long strings for Excel safety."""
    if pd.isna(val) or val is None:
        return val
    s = ILLEGAL_CHARACTERS_RE.sub("", str(val))
    if len(s) > 32760:
        s = s[:32757] + "..."
    return s


def style_worksheet(ws, tab_color="1F4E79", freeze="A2"):
    """Apply professional corporate styling to openpyxl worksheet."""
    ws.sheet_properties.tabColor = tab_color
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color=tab_color, end_color=tab_color, fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=False)

    thin_border = Border(
        bottom=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="EEEEEE")
    )

    # Style header row
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.row_dimensions[1].height = 28

    # Style data rows and column widths
    col_widths = {
        "split": 10,
        "openparldata_id": 16,
        "affair_external_id": 18,
        "number": 12,
        "type_en": 22,
        "type_de": 26,
        "begin_date": 14,
        "title_de": 45,
        "title_fr": 45,
        "title_it": 45,
        "submitted_text_de": 60,
        "submitted_text_fr": 60,
        "submitted_text_it": 60,
        "reasoning_de": 60,
        "reasoning_fr": 60,
        "reasoning_it": 60,
        "answer_de": 60,
        "answer_fr": 60,
        "answer_it": 60,
        "topic_ids": 16,
        "topics_de": 32,
        "topics_fr": 32,
        "topics_it": 32,
        "primary_topic": 14,
        "descriptors_de": 35,
    }

    for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
        col_name = col[0].value
        col_letter = get_column_letter(col_idx)
        width = col_widths.get(col_name, 20)
        ws.column_dimensions[col_letter].width = width

    # Subtle bottom border on all rows
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border


def build_summary_sheet(wb, combined_df, train_df, val_df, test_df, text_cols):
    """Create a formatted Summary_Stats worksheet with metrics and tables."""
    ws = wb.create_sheet(title="Summary_Stats")
    ws.sheet_properties.tabColor = "2C3E50"

    title_font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="555555")
    section_font = Font(name="Calibri", size=13, bold=True, color="1F4E79")
    tbl_header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    tbl_header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True)
    total_fill = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    # Title block
    ws["A1"] = "OpenParl-Apertus Core Dataset Summary"
    ws["A1"].font = title_font
    ws["A2"] = "Swiss Parliamentary Affairs (Curia Vista / OpenParlData.ch) — 2,000 Sample Core Corpus"
    ws["A2"].font = subtitle_font

    current_row = 4

    # ── Table 1: Split Breakdown ─────────────────────────────────────────────
    ws.cell(row=current_row, column=1, value="1. Dataset Split Overview").font = section_font
    current_row += 1

    headers = ["Split", "Samples", "Percentage", "Target Use Case"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=h)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    current_row += 1

    split_rows = [
        ("Train", len(train_df), f"{len(train_df)/len(combined_df):.1%}", "Model fine-tuning & prompt adaptation (70%)"),
        ("Validation", len(val_df), f"{len(val_df)/len(combined_df):.1%}", "Hyperparameter selection & checkpoint validation (15%)"),
        ("Test", len(test_df), f"{len(test_df)/len(combined_df):.1%}", "Unseen evaluation & final metric reporting (15%)"),
    ]
    for r in split_rows:
        ws.cell(row=current_row, column=1, value=r[0]).alignment = Alignment(horizontal="left")
        ws.cell(row=current_row, column=2, value=r[1]).alignment = Alignment(horizontal="right")
        ws.cell(row=current_row, column=3, value=r[2]).alignment = Alignment(horizontal="right")
        ws.cell(row=current_row, column=4, value=r[3]).alignment = Alignment(horizontal="left")
        for c in range(1, 5):
            ws.cell(row=current_row, column=c).border = thin_border
        current_row += 1

    # Total row
    ws.cell(row=current_row, column=1, value="Total Core Dataset").font = total_font
    ws.cell(row=current_row, column=2, value=len(combined_df)).font = total_font
    ws.cell(row=current_row, column=3, value="100.0%").font = total_font
    ws.cell(row=current_row, column=4, value="Full gold-standard classification sample").font = total_font
    for c in range(1, 5):
        ws.cell(row=current_row, column=c).fill = total_fill
        ws.cell(row=current_row, column=c).border = thin_border
        if c in (2, 3):
            ws.cell(row=current_row, column=c).alignment = Alignment(horizontal="right")
    current_row += 3

    # ── Table 2: Primary Topic Distribution ──────────────────────────────────
    ws.cell(row=current_row, column=1, value="2. Primary Topic Distribution Across Splits").font = section_font
    current_row += 1

    topic_headers = [
        "Primary Topic ID", "Topic (German)", "Topic (French)", "Topic (Italian)",
        "Train Count", "Val Count", "Test Count", "Total Count", "Dataset Share (%)"
    ]
    for col_idx, h in enumerate(topic_headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=h)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    current_row += 1

    # Aggregate topic counts
    topic_summary = combined_df.groupby("primary_topic").agg(
        total_count=("openparldata_id", "count"),
        topic_de=("topics_de", lambda s: str(s.iloc[0]).split(" | ")[0] if pd.notna(s.iloc[0]) else ""),
        topic_fr=("topics_fr", lambda s: str(s.iloc[0]).split(" | ")[0] if pd.notna(s.iloc[0]) else ""),
        topic_it=("topics_it", lambda s: str(s.iloc[0]).split(" | ")[0] if pd.notna(s.iloc[0]) else ""),
    ).sort_values("total_count", ascending=False)

    train_counts = train_df["primary_topic"].value_counts()
    val_counts = val_df["primary_topic"].value_counts()
    test_counts = test_df["primary_topic"].value_counts()

    for tid, row_data in topic_summary.iterrows():
        n_tot = row_data["total_count"]
        n_tr = train_counts.get(tid, 0)
        n_va = val_counts.get(tid, 0)
        n_te = test_counts.get(tid, 0)
        share = n_tot / len(combined_df)

        ws.cell(row=current_row, column=1, value=str(tid)).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=2, value=row_data["topic_de"])
        ws.cell(row=current_row, column=3, value=row_data["topic_fr"])
        ws.cell(row=current_row, column=4, value=row_data["topic_it"])
        ws.cell(row=current_row, column=5, value=n_tr).alignment = Alignment(horizontal="right")
        ws.cell(row=current_row, column=6, value=n_va).alignment = Alignment(horizontal="right")
        ws.cell(row=current_row, column=7, value=n_te).alignment = Alignment(horizontal="right")
        ws.cell(row=current_row, column=8, value=n_tot).alignment = Alignment(horizontal="right")
        ws.cell(row=current_row, column=9, value=f"{share:.2%}").alignment = Alignment(horizontal="right")

        for c in range(1, 10):
            ws.cell(row=current_row, column=c).border = thin_border
        current_row += 1

    current_row += 2

    # ── Table 3: Text Field Availability ─────────────────────────────────────
    ws.cell(row=current_row, column=1, value="3. Multilingual Text Completeness").font = section_font
    current_row += 1

    text_headers = ["Field Column", "Text Type", "Language", "Available (Count)", "Completeness (%)"]
    for col_idx, h in enumerate(text_headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=h)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    current_row += 1

    lang_map = {"de": "German (DE)", "fr": "French (FR)", "it": "Italian (IT)"}
    for col in text_cols:
        parts = col.split("_")
        prefix = "_".join(parts[:-1])
        lang = parts[-1]
        cnt = combined_df[col].notna().sum()
        pct = cnt / len(combined_df)

        ws.cell(row=current_row, column=1, value=col)
        ws.cell(row=current_row, column=2, value=prefix)
        ws.cell(row=current_row, column=3, value=lang_map.get(lang, lang.upper()))
        ws.cell(row=current_row, column=4, value=cnt).alignment = Alignment(horizontal="right")
        ws.cell(row=current_row, column=5, value=f"{pct:.1%}").alignment = Alignment(horizontal="right")

        for c in range(1, 6):
            ws.cell(row=current_row, column=c).border = thin_border
        current_row += 1

    # Adjust column widths on Summary sheet
    summary_widths = {1: 22, 2: 28, 3: 28, 4: 28, 5: 14, 6: 14, 7: 14, 8: 14, 9: 18}
    for col_idx, w in summary_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def main():
    parser = argparse.ArgumentParser(description="Generate parliamentary affair topic classification dataset.")
    parser.add_argument("--samples", "-n", type=int, default=DEFAULT_SAMPLES, help="Total number of samples (default: 2000)")
    parser.add_argument("--seed", "-s", type=int, default=DEFAULT_SEED, help="Random seed for reproducibility (default: 42)")
    args = parser.parse_args()

    n_samples = args.samples
    seed = args.seed

    print("=" * 70)
    print(f"Creating Core Dataset: {n_samples:,} entries (Seed: {seed})")
    print("=" * 70)

    # ── 1. Load affair metadata + labels ────────────────────────────────────────
    print("1. Loading affair metadata and labels...")
    keywords_df = pd.read_csv(
        KEYWORDS_CSV,
        dtype={"affair_external_id": str, "number": str, "topic_ids": str},
        low_memory=False,
    )
    # Keep only affairs that have topic labels
    keywords_df = keywords_df[keywords_df["topic_ids"].notna()].copy()
    valid_ids = set(keywords_df["openparldata_id"])
    print(f"   → {len(keywords_df):,} affairs with topic labels")

    # ── 2. Load texts ────────────────────────────────────────────────────────────
    print("2. Loading affair texts (che_affair_texts.ndjson.gz)...")
    texts_df = pd.read_json(TEXTS_NDJSON, lines=True)
    print(f"   → {len(texts_df):,} total text rows loaded")

    # Filter early to valid IDs to accelerate subsequent processing
    texts_df = texts_df[texts_df["openparldata_id"].isin(valid_ids)].copy()
    print(f"   → {len(texts_df):,} texts matched to labelled affairs")

    # ── 3. Strip HTML from text fields ──────────────────────────────────────────
    print("3. Stripping HTML tags and cleaning text fields...")
    TAG_RE = re.compile(r"<[^>]+>")

    def strip_html(text):
        if pd.isna(text):
            return None
        text = TAG_RE.sub(" ", str(text))
        text = re.sub(r"\s+", " ", text).strip()
        return text if text else None

    for col in ["text_de", "text_fr", "text_it"]:
        texts_df[col] = texts_df[col].apply(strip_html)

    # ── 4. Map type_de to canonical column names ────────────────────────────────
    print("4. Mapping text types to canonical names...")
    TYPE_MAP = {
        "Titel des Geschäftes": "title",
        "Eingereichter Text": "submitted_text",
        "Begründung": "reasoning",
        "Antwort BR / Büro": "answer",
    }
    texts_df["text_type"] = texts_df["type_de"].map(TYPE_MAP)
    texts_df = texts_df[texts_df["text_type"].notna()].copy()
    print(f"   → {len(texts_df):,} texts after filtering to main 4 types")

    # ── 5. Vectorized Pivot ─────────────────────────────────────────────────────
    print("5. Pivoting texts into one row per affair...")
    # For affairs with multiple texts of the same type, keep the first
    texts_df = texts_df.sort_values(["openparldata_id", "text_id"])
    texts_df = texts_df.drop_duplicates(subset=["openparldata_id", "text_type"], keep="first")

    piv = texts_df.pivot(index="openparldata_id", columns="text_type", values=["text_de", "text_fr", "text_it"])
    piv.columns = [f"{text_type}_{lang.replace('text_', '')}" for lang, text_type in piv.columns]
    texts_wide = piv.reset_index()
    print(f"   → {len(texts_wide):,} affairs with pivoted multilingual texts")

    # ── 6. Merge texts with labels ──────────────────────────────────────────────
    print("6. Merging texts with metadata and labels...")
    merged = pd.merge(
        keywords_df[[
            "openparldata_id", "affair_external_id", "number",
            "type_en", "type_de", "begin_date",
            "topic_ids", "topics_de", "topics_fr", "topics_it",
            "descriptors_de",
        ]],
        texts_wide,
        on="openparldata_id",
        how="inner",
    )
    print(f"   → {len(merged):,} affairs after inner join")

    # ── 7. Filter: require at least submitted_text_de ───────────────────────────
    print("7. Filtering to affairs with at least German submitted text...")
    merged = merged[merged["submitted_text_de"].notna()].copy()
    print(f"   → {len(merged):,} affairs available for sampling")

    # ── 8. Stratified Sampling ──────────────────────────────────────────────────
    print(f"8. Sampling {n_samples:,} affairs with topic-stratified sampling...")
    merged["primary_topic"] = merged["topic_ids"].str.split(r"\s*\|\s*").str[0]

    topic_counts = merged["primary_topic"].value_counts()
    eligible_topics = topic_counts[topic_counts >= 3].index
    merged_eligible = merged[merged["primary_topic"].isin(eligible_topics)]
    print(f"   → {len(merged_eligible):,} affairs from topics with >= 3 examples ({len(eligible_topics)} topics)")

    rng = np.random.RandomState(seed)
    sampled = merged_eligible.groupby("primary_topic", group_keys=False).apply(
        lambda g: g.sample(
            n=max(1, int(round(n_samples * len(g) / len(merged_eligible)))),
            random_state=rng,
        )
    )

    # Trim or pad to exactly n_samples
    if len(sampled) > n_samples:
        sampled = sampled.sample(n=n_samples, random_state=seed)
    elif len(sampled) < n_samples:
        remaining = merged_eligible[~merged_eligible.index.isin(sampled.index)]
        extra = remaining.sample(n=n_samples - len(sampled), random_state=seed)
        sampled = pd.concat([sampled, extra])

    sampled = sampled.sample(frac=1, random_state=seed).reset_index(drop=True)
    print(f"   → Successfully sampled exactly {len(sampled):,} affairs")

    # ── 9. Split into train / val / test (70% / 15% / 15%) ──────────────────────
    print("9. Splitting into Train (70%), Validation (15%), Test (15%)...")
    n_train = int(round(0.70 * n_samples))
    n_val = int(round(0.15 * n_samples))
    n_test = n_samples - n_train - n_val

    train_df = sampled.iloc[:n_train].copy()
    val_df = sampled.iloc[n_train: n_train + n_val].copy()
    test_df = sampled.iloc[n_train + n_val:].copy()

    train_df["split"] = "train"
    val_df["split"] = "val"
    test_df["split"] = "test"

    combined = pd.concat([train_df, val_df, test_df]).reset_index(drop=True)

    # ── 10. Column Ordering ─────────────────────────────────────────────────────
    text_cols = []
    for prefix in ["title", "submitted_text", "reasoning", "answer"]:
        for lang in ["de", "fr", "it"]:
            col = f"{prefix}_{lang}"
            if col in sampled.columns:
                text_cols.append(col)

    meta_cols = [
        "openparldata_id", "affair_external_id", "number",
        "type_en", "type_de", "begin_date",
    ]
    label_cols = [
        "topic_ids", "topics_de", "topics_fr", "topics_it",
        "primary_topic", "descriptors_de",
    ]
    final_cols = ["split"] + meta_cols + text_cols + label_cols

    train_df = train_df[final_cols]
    val_df = val_df[final_cols]
    test_df = test_df[final_cols]
    combined = combined[final_cols]

    # ── 11. Save CSVs ────────────────────────────────────────────────────────────
    print("10. Writing CSV splits...")
    for name, df in [("train", train_df), ("val", val_df), ("test", test_df)]:
        out_csv = OUTPUT_DIR / f"{name}.csv"
        df.to_csv(out_csv, index=False, encoding="utf-8")
        print(f"    ✓ {name}.csv ({len(df):,} rows) → {out_csv}")

    combined_csv = OUTPUT_DIR / f"combined_{n_samples}.csv"
    combined.to_csv(combined_csv, index=False, encoding="utf-8")
    print(f"    ✓ combined_{n_samples}.csv ({len(combined):,} rows) → {combined_csv}")

    # ── 12. Save Single Excel File with Tabs ─────────────────────────────────────
    print(f"11. Generating single Excel workbook: {EXCEL_OUTPUT_PATH}...")
    # Sanitize dataframe copies for Excel (prevent illegal chars & cell overflow)
    excel_combined = combined.copy()
    excel_train = train_df.copy()
    excel_val = val_df.copy()
    excel_test = test_df.copy()

    for col in excel_combined.select_dtypes(include=["object"]).columns:
        excel_combined[col] = excel_combined[col].apply(sanitize_text_for_excel)
        excel_train[col] = excel_train[col].apply(sanitize_text_for_excel)
        excel_val[col] = excel_val[col].apply(sanitize_text_for_excel)
        excel_test[col] = excel_test[col].apply(sanitize_text_for_excel)

    with pd.ExcelWriter(EXCEL_OUTPUT_PATH, engine="openpyxl") as writer:
        # Sheet 1: Core_Dataset (all 2000 entries)
        excel_combined.to_excel(writer, sheet_name="Core_Dataset", index=False)
        style_worksheet(writer.sheets["Core_Dataset"], tab_color="1F4E79")

        # Sheet 2: Train (1400 entries)
        excel_train.to_excel(writer, sheet_name="Train", index=False)
        style_worksheet(writer.sheets["Train"], tab_color="2E7D32")

        # Sheet 3: Validation (300 entries)
        excel_val.to_excel(writer, sheet_name="Validation", index=False)
        style_worksheet(writer.sheets["Validation"], tab_color="E65100")

        # Sheet 4: Test (300 entries)
        excel_test.to_excel(writer, sheet_name="Test", index=False)
        style_worksheet(writer.sheets["Test"], tab_color="C2185B")

        # Sheet 5: Summary & Metadata
        build_summary_sheet(writer.book, combined, train_df, val_df, test_df, text_cols)

    # Also save/copy to data/splits/ for convenience
    import shutil
    shutil.copy2(EXCEL_OUTPUT_PATH, EXCEL_SPLITS_PATH)
    print(f"    ✓ Primary Excel Workbook created: {EXCEL_OUTPUT_PATH}")
    print(f"    ✓ Copy also available in splits/: {EXCEL_SPLITS_PATH}")

    # ── 13. Summary Statistics ──────────────────────────────────────────────────
    print("\n" + "=" * 70)
    print("DATASET GENERATION SUMMARY")
    print("=" * 70)
    print(f"  Total samples in Core Dataset: {len(combined):,}")
    print(f"  Train samples (70%):           {len(train_df):,}")
    print(f"  Validation samples (15%):      {len(val_df):,}")
    print(f"  Test samples (15%):            {len(test_df):,}")
    print(f"  Unique primary topics:         {combined['primary_topic'].nunique()}")
    multi_label = combined["topic_ids"].str.contains(r"\|", na=False).sum()
    print(f"  Multi-label affairs:           {multi_label:,} ({multi_label / len(combined):.1%})")

    print("\n  Top 10 Primary Topics:")
    for tid, count in combined["primary_topic"].value_counts().head(10).items():
        sample_title = combined[combined["primary_topic"] == tid]["topics_de"].iloc[0].split(" | ")[0]
        print(f"    - [{tid:>4s}] {sample_title:35s}: {count:4d} samples ({count/len(combined):.1%})")

    print("\n  Text Column Availability:")
    for col in text_cols:
        pct = combined[col].notna().mean() * 100
        print(f"    {col:22s}: {pct:5.1f}% non-null ({combined[col].notna().sum():,} rows)")

    print(f"\nExcel File:  {EXCEL_OUTPUT_PATH}")
    print(f"CSV Outputs: {OUTPUT_DIR}")


if __name__ == "__main__":
    main()
